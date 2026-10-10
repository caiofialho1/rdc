"""Verificação local: python tests/test_painel.py (Playwright + Edge).

Todas as chamadas externas são interceptadas. Nenhum dado real é consultado.
"""
import datetime
import json
import tempfile
from pathlib import Path
import openpyxl
from playwright.sync_api import sync_playwright, expect

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(tempfile.gettempdir()) / 'rdc-painel-review'
OUT.mkdir(exist_ok=True)


def planilha_mc(caminho, ativo, itens, contrato='ROTINA'):
    """MC_PAC mínima: rótulos nas linhas 3-7, cabeçalho na 9, itens a partir da 10 e tabela diária ao lado."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'MC_PAC'
    ws['B3'], ws['D3'] = 'Cliente:', 'Cliente de teste'
    ws['B4'], ws['D4'] = 'Contrato:', contrato
    ws['B7'], ws['D7'] = 'Ativo:', ativo
    ws['Y5'], ws['Z5'] = 'Início', datetime.datetime(2026, 8, 21)
    ws['Y6'], ws['Z6'] = 'Término', datetime.datetime(2026, 9, 20)
    cab = ['Item', 'Subitem', 'Referência Projeto', 'Ordem de Manutenção', 'Centro de Custo', 'Área de Referência',
           'Descrição da Estrutura/Perfil', 'Larg/Per. (m)', 'Comp. (m)', 'Quantidade', 'Lados', 'Critério', 'Prev.Total (m²)',
           'Exec Atual (%)', 'Exec. Atual (m²)', 'Observações']
    for i, h in enumerate(cab):
        ws.cell(9, 2 + i, h)
    for j, e in enumerate(['SLV', 'A05']):
        ws.cell(9, 19 + j, e)             # etapas por item
    ws.cell(9, 24, 'DATA')
    for j, e in enumerate(['SLV', 'A05']):
        ws.cell(9, 25 + j, e)             # tabela diária
    for k, (desc, prev, pct, dt) in enumerate(itens):
        r = 10 + k
        for c, v in zip(range(2, 18), [k + 1, k + 1, 'N/D', 202500000001, 1390337, 'ÁREA', desc, 1, 1, 1, 1, 1, prev, pct, prev * pct, '-']):
            ws.cell(r, c, v)
        ws.cell(r, 19, dt).number_format = 'dd/mm/yyyy'
        ws.cell(r, 20, 'A05')
        ws.cell(r, 24, dt).number_format = 'dd/mm/yyyy'
        ws.cell(r, 25, prev * pct)
    wb.save(caminho)


def fixture():
    rdcs, atividades = [], []
    for i, (status, pts, espera) in enumerate([
        ('Enviado', True, 30), ('Validado', True, None),
        ('Validado', False, None), ('Cancelado', False, None)
    ]):
        codigo = f'TESTE-{i}'
        rdcs.append(dict(codigo=codigo, data_rdc=f'2026-10-0{i+1}',
                         responsavel='Responsável de teste', area='Área teste',
                         ativo='AT-01', om='1234', status=status,
                         contrato=None if i == 0 else '5900108690',
                         contrato_nome='ROTINA', pts=pts, espera_pts_min=espera,
                         bloqueio=False, hh_min=120, hh_impacto_min=60,
                         min_impacto=30, min_atividades=60, efetivo_max=2,
                         qtd_atividades=1, qtd_impactos=1, qtd_fotos=0))
        atividades.append(dict(codigo=f'A-{i}', rdc=codigo, ordem=1,
                               disciplina='PAC' if i == 0 else 'REC',
                               servico='Teste', efetivo_total=2, duracao_min=60))
    return dict(ok=True, rdcs=rdcs, atividades=atividades, fotos=[], impactos=[
        dict(codigo=f'I-{i}', rdc=r['codigo'], ordem=1, descricao='Impacto de teste',
             duracao_min=30) for i, r in enumerate(rdcs)
    ])


with sync_playwright() as p:
    browser = p.chromium.launch(channel='msedge', headless=True)
    context = browser.new_context(viewport={'width': 1440, 'height': 1000},
                                  timezone_id='America/Sao_Paulo')
    state = {'mode': 'ok'}
    export_calls = []
    prod_calls = []
    resp_calls = []
    carga_banco, carga_enviada = [], []

    def route(req):
        if req.request.url.endswith('/rdc_admin_login'):
            payload = {'ok': True, 'nome': 'Teste local'}
        elif req.request.url.endswith('/rdc_admin_exportar'):
            args = json.loads(req.request.post_data or '{}')
            export_calls.append(args)
            if state['mode'] == 'network':
                req.abort('failed')
                return
            payload = ({'ok': False, 'erro': 'Falha de teste'} if state['mode'] == 'error'
                       or state['mode'] == 'cycle_error' and args.get('p_ini') == '2026-09-21'
                       else dict(ok=True, rdcs=[], atividades=[], impactos=[], fotos=[])
                       if state['mode'] == 'empty' else fixture())
        elif req.request.url.endswith('/rdc_admin_producao'):
            prod_calls.append(json.loads(req.request.post_data or '{}'))
            payload = dict(ok=True, linhas=[
                dict(tag_chave='PO840KPE', data='2026-10-01', om='202502067677', disciplina='PAC',
                     situacao='Produção com RDC', ativos_mc=['PO_840K_PE (PISO 2)'], rdcs=['TESTE-0'],
                     horas=8, homem_hora=40, slv=691.69, a05=625.05),
                dict(tag_chave='PO840KPE', data='2026-10-02', om='202502067677', disciplina='PAC',
                     situacao='Produção sem RDC', ativos_mc=['PO_840K_PE (PISO 2)'], slv=12.5),
                dict(tag_chave='TR811K02', data='2026-10-03', om='9', disciplina='REC',
                     situacao='RDC sem produção', tags_rdc=['TR_811K_02'], rdcs=['TESTE-2'], horas=4, homem_hora=8)])
        elif req.request.url.endswith('/rdc_admin_mc_status'):
            payload = dict(ok=True, arquivos=carga_banco)
        elif req.request.url.endswith('/rdc_admin_mc_importar'):
            arq = json.loads(req.request.post_data or '{}')['p_arquivo']
            carga_enviada.append(arq)
            payload = dict(ok=True, itens=len(arq.get('itens', [])), diario=len(arq.get('diario', [])))
        elif req.request.url.endswith('/rdc_admin_producao_resp'):
            resp_calls.append(json.loads(req.request.post_data or '{}'))
            base = dict(tag='PO_840K_PE', data='2026-10-01', om='1', disciplina='PAC', unidade='m2',
                        situacao='Produção com RDC', etapas_sem_preco=None)
            payload = dict(ok=True, precos=[], linhas=[
                dict(base, responsavel='Ana', slv=60.0, a05=30.0, valor=1000.0, homem_hora=30, horas=6, participacao=0.6),
                dict(base, responsavel='Bruno', slv=40.0, a05=20.0, valor=500.0, homem_hora=20, horas=4, participacao=0.4),
                dict(base, responsavel='Sem RDC', data='2026-10-02', situacao='Produção sem RDC', slv=10.0, valor=250.0, etapas_sem_preco='ST2')])
        else:
            req.abort('blockedbyclient')
            return
        req.fulfill(status=200, content_type='application/json', body=json.dumps(payload))

    context.route('https://**/*', route)
    page = context.new_page()
    errors = []
    page.on('pageerror', lambda error: errors.append(str(error)))
    page.goto((ROOT / 'painel.html').as_uri())
    expect(page.locator('#loginTitulo')).to_have_text('Painel de campo')
    expect(page.locator('.login-credit')).to_have_text('Desenvolvido por Fialho Consultoria de Planejamento LTDA.')
    assert page.evaluate("""() => {
      const card = document.querySelector('#scrLogin').getBoundingClientRect();
      const logo = document.querySelector('.login-logo img');
      return card.width <= 440 && card.width >= 400 && logo.complete && logo.naturalWidth > 0
        && document.querySelector('.login-logo').getBoundingClientRect().width === 136;
    }""")
    page.screenshot(path=str(OUT / 'login-1440.png'), full_page=True)
    page.set_viewport_size({'width': 390, 'height': 844})
    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    assert page.locator('#scrLogin').evaluate('e => e.getBoundingClientRect().width <= innerWidth - 32')
    assert page.locator('.login-logo').evaluate('e => e.getBoundingClientRect().width === 116')
    page.screenshot(path=str(OUT / 'login-390.png'), full_page=True)
    page.set_viewport_size({'width': 1440, 'height': 1000})
    expect(page.locator('#btnExcel')).to_be_disabled()
    page.locator('#btnEntrar').click()
    expect(page.locator('#lErr')).to_have_text('Informe nome e senha.')
    page.locator('#lNome').fill('Teste local')
    page.locator('#lSenha').fill('teste-simulado')
    page.locator('#lSenha').press('Enter')
    expect(page.locator('#kpis')).to_contain_text('6')
    shapes = page.evaluate("""() => {
      const radius = selector => getComputedStyle(document.querySelector(selector)).borderTopLeftRadius;
      return Object.fromEntries(['.top-btn', '.button', '.ctl', '.search-wrap', '.period-group button', '.metric', '.bloco', '.tbl-wrap', '.filter-panel', '.login', '.section-number'].map(s => [s, radius(s)]));
    }""")
    assert all(shapes[s] == '4px' for s in ['.top-btn', '.button', '.ctl', '.search-wrap', '.period-group button']), shapes
    assert all(shapes[s] == '6px' for s in ['.metric', '.bloco', '.tbl-wrap', '.filter-panel', '.login']), shapes
    assert shapes['.section-number'] == '3px', shapes
    assert page.evaluate("""() => {
      const height = selector => document.querySelector(selector).getBoundingClientRect().height;
      const headings = ['#overview', '#sec-analise>summary', '#sec-emissao>summary'].map(height);
      return headings.every(h => h >= 40 && h <= 48) && Math.max(...headings) - Math.min(...headings) <= 6;
    }""")
    assert page.evaluate("""() => {
      const color = selector => getComputedStyle(document.querySelector(selector)).backgroundColor;
      const levels = ['#overview', '#sec-analise>summary', '.disc-card h3'].map(color);
      const brightness = value => value.match(/[0-9]+/g).slice(0, 3).map(Number).reduce((a, b) => a + b);
      return new Set(levels).size === 3 && levels.every((value, index) =>
        index === 0 || brightness(levels[index - 1]) < brightness(value));
    }""")
    page.evaluate("$('fIni').value='2026-10-01'; $('fFim').value='2026-10-07'; $('fFim').dispatchEvent(new Event('change'))")
    page.locator('#btnAtualizar').click()
    expect(page.locator('#perCap')).to_contain_text('07/10/2026')
    assert page.locator('#matriz thead tr.dia th').count() == 8  # nome + 7 dias do filtro passado
    assert page.locator('#matriz .mat-stage.com-marcos').count() == 0
    expect(page.locator('.production-hero .metric-value')).to_have_text('6HH')
    expect(page.locator('#resumoTexto')).to_contain_text('1 RDC aguardando validação')
    expect(page.locator('#resumoTexto')).to_contain_text('1 sem contrato')
    expect(page.locator('#resumoTexto')).to_contain_text('1 com PTS sem tempo')
    assert page.locator('#resumoTexto strong').all_text_contents() == ['1', '1', '1']
    assert page.locator('#resumoTexto .summary-item .ic').count() == 3
    assert page.locator('#resumoTexto .summary-item.attention').count() == 3
    expect(page.locator('#kpis .metric').nth(4)).to_have_attribute('title', 'Média · 1 de 2 RDC com PTS')
    assert page.locator('#kpis .metric-description, .bloco>summary p, .bloco>summary .meta, .page-heading p').count() == 0
    page.locator('#sec-emissao').evaluate('(e)=>e.open=true')
    page.locator('#sec-resp').evaluate('(e)=>e.open=true')
    expect(page.locator('#gResp')).to_have_count(0)
    expect(page.locator('#tblResp tbody tr.g-row')).to_have_count(0)
    expect(page.locator('#graficosResponsaveis .resp-chart')).to_have_count(1)
    expect(page.locator('#graficosResponsaveis .resp-echart svg')).to_have_count(1)
    expect(page.locator('#graficosResponsaveis .resp-echart-donut')).to_have_count(0)
    assert page.evaluate("""() => {
      const table = document.querySelector('#tblResp').getBoundingClientRect();
      const charts = document.querySelector('.resp-analytics').getBoundingClientRect();
      const option = echarts.getInstanceByDom(document.querySelector('.resp-echart')).getOption();
      const sum = values => values.reduce((total, value) => total + value, 0);
      return charts.left >= table.right
        && option.xAxis[0].type === 'value' && option.xAxis[0].min === 0
        && option.yAxis[0].type === 'category'
        && option.series.every(serie => serie.type === 'bar' && serie.stack === 'hh-responsavel')
        && option.series[0].data.length === 1
        && Math.abs(sum(option.series[0].data) - 6) < 0.001
        && Math.abs(sum(option.series[1].data) - 3) < 0.001;
    }""")
    page.evaluate("""() => PainelCharts.renderResponsaveis([
      {nome:'Valdeci Manga Rosa',hh:36000,hhimp:300,rdc:6,aguard:1},
      {nome:'Carlos Caixão',hh:19500,hhimp:0,rdc:8,aguard:0},
      {nome:'Jamilson Santos',hh:13200,hhimp:1800,rdc:4,aguard:1},
      {nome:'Leandro Costa',hh:4500,hhimp:300,rdc:2,aguard:0},
      {nome:'Dorivan Costa',hh:2700,hhimp:0,rdc:3,aguard:1}
    ])""")
    page.wait_for_function("echarts.getInstanceByDom(document.querySelector('.resp-echart')).getOption().yAxis[0].data.length === 5")
    assert page.evaluate("""() => {
      const chart = echarts.getInstanceByDom(document.querySelector('.resp-echart')).getOption();
      return chart.series.every(serie => serie.type === 'bar' && serie.stack === 'hh-responsavel' && serie.data.length === 5)
        && chart.xAxis[0].type === 'value' && chart.yAxis[0].type === 'category';
    }""")
    page.locator('.resp-analytics').screenshot(path=str(OUT / 'comparativo-responsaveis-1440.png'))
    page.set_viewport_size({'width': 390, 'height': 844})
    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    page.locator('.resp-analytics').screenshot(path=str(OUT / 'comparativo-responsaveis-390.png'))
    page.set_viewport_size({'width': 1440, 'height': 1000})
    page.evaluate('render()')
    dimensions = page.evaluate("""() => ({
      matrix: document.querySelector('#matriz tbody tr').getBoundingClientRect().height,
      table: document.querySelector('#tblResp tbody tr').getBoundingClientRect().height,
      avatar: document.querySelector('#tblResp .resp-avatar').getBoundingClientRect().height,
      matrixAvatar: document.querySelector('#matriz .resp-avatar').getBoundingClientRect().height
    })""")
    assert dimensions == {'matrix': 49, 'table': 47, 'avatar': 46, 'matrixAvatar': 48}, dimensions
    assert page.evaluate("""() => {
      const cells = [...document.querySelectorAll('#tblResp thead th.num')];
      const widths = cells.map(e => e.getBoundingClientRect().width);
      const headersFit = cells.every(e => e.scrollWidth <= e.clientWidth + 1);
      return widths.length === 7 && widths.every(w => Math.abs(w - 100) <= 1) && headersFit;
    }""")
    page.locator('#sec-imp').evaluate('(e)=>e.open=true')
    page.locator('#sec-rdc').evaluate('(e)=>e.open=true')
    assert page.evaluate("""() => ['#tblImp', '#tblRdc'].every(selector => {
      const cells = [...document.querySelectorAll(`${selector} thead th.num`)];
      return cells.length > 0 && cells.every(e => Math.abs(e.getBoundingClientRect().width - 100) <= 1);
    })""")
    page.locator('#sec-imp').evaluate('(e)=>e.open=false')
    page.locator('#sec-rdc').evaluate('(e)=>e.open=false')
    page.locator('#sec-prod').evaluate('(e)=>e.open=true')
    expect(page.locator('#tblProd tbody tr')).to_have_count(3)
    expect(page.locator('#tblProd tbody tr').nth(0)).to_contain_text('691,69')
    expect(page.locator('#tblProd tbody tr').nth(1)).to_contain_text('Produção sem RDC')
    expect(page.locator('#prodResumo')).to_contain_text('1 com RDC · 1 produção sem RDC · 1 RDC sem produção')
    assert len(prod_calls) == 1 and prod_calls[0]['p_nome'] and prod_calls[0]['p_ini'] <= prod_calls[0]['p_fim'], prod_calls
    page.locator('#sec-prod').evaluate('(e)=>e.open=false')
    page.locator('#sec-presp').evaluate('(e)=>e.open=true')
    expect(page.locator('#tblProdResp tr.g-row')).to_have_count(3)
    expect(page.locator('#tblProdResp')).to_contain_text('Ana')
    expect(page.locator('#prodRespResumo')).to_contain_text('R$ 1.750,00')
    expect(page.locator('#prodRespResumo')).to_contain_text('14,3% sem RDC')
    expect(page.locator('#prodRespResumo')).to_contain_text('PAC ST2')
    page.locator('#gProdResp').select_option('disciplina')
    expect(page.locator('#tblProdResp tr.g-row')).to_have_count(1)
    assert len(resp_calls) == 1 and resp_calls[0]['p_ini'] == '2020-01-01', resp_calls  # padrão: todas as datas
    assert prod_calls[0]['p_ini'] == '2020-01-01', prod_calls
    page.locator('#escopoResp').select_option('painel')
    expect(page.locator('#escopoProd')).to_have_value('painel')  # os dois seletores andam juntos
    page.wait_for_timeout(300)
    assert len(resp_calls) == 2 and resp_calls[1]['p_ini'] != '2020-01-01', resp_calls
    page.locator('#escopoResp').select_option('todas')
    page.wait_for_timeout(300)
    page.locator('#gProdResp').select_option('responsavel')
    page.locator('#sec-presp').screenshot(path=str(OUT / 'producao-responsavel-1440.png'))
    page.set_viewport_size({'width': 390, 'height': 900})
    page.locator('#sec-presp').screenshot(path=str(OUT / 'producao-responsavel-390.png'))
    page.set_viewport_size({'width': 1440, 'height': 1000})
    page.locator('#sec-presp').evaluate('(e)=>e.open=false')
    # botão "Atualizar carga produção": lê a pasta no navegador e envia só o que mudou
    pasta_mc = Path(tempfile.mkdtemp(prefix='rdc-mc-'))
    (pasta_mc / 'TAC').mkdir()
    (pasta_mc / 'TELHADO').mkdir()
    planilha_mc(pasta_mc / 'TAC' / 'MC_ORC - TESTE.xlsx', 'PO_840K_PE (PISO 2)',
                [('VIGA A', 10.0, 0.5, datetime.datetime(2026, 9, 1)), ('VIGA B', 20.0, 1.0, datetime.datetime(2026, 9, 2))])
    planilha_mc(pasta_mc / 'TELHADO' / 'MC_ORC - TELHADO.xlsx', 'PO_813K_RI', [('TELHA', 5.0, 1.0, datetime.datetime(2026, 9, 3))], 'TELHADO')
    expect(page.locator('#btnCarga')).to_have_text('Atualizar carga produção')
    page.locator('#cargaPasta').set_input_files(str(pasta_mc))
    expect(page.locator('#cargaPasso')).to_contain_text('1 planilha(s) gravada(s) (2 itens)', timeout=30000)
    expect(page.locator('#cargaLog')).to_contain_text('1 arquivo(s) da pasta TELHADO não foram lidos')
    assert len(carga_enviada) == 1, carga_enviada
    f = carga_enviada[0]
    assert f['chave'] == 'MC_ORC - TESTE.xlsx|MC_PAC' and f['contrato_codigo'] == '5900108690' and f['ativo'] == 'PO_840K_PE (PISO 2)', f
    assert f['periodo_ini'] == '2026-08-21' and f['periodo_fim'] == '2026-09-20' and f['disciplina'] == 'PAC' and f['unidade'] == 'm2', f
    assert [i['etapas'] for i in f['itens']] == [{'SLV': '2026-09-01'}, {'SLV': '2026-09-02'}], f['itens']
    assert [round(i['exec_qtd'], 6) for i in f['itens']] == [5.0, 20.0] and f['itens'][0]['om'] == '202500000001', f['itens']
    assert sorted((d['data'], d['etapa'], d['valor']) for d in f['diario']) == [('2026-09-01', 'SLV', 5.0), ('2026-09-02', 'SLV', 20.0)], f['diario']
    carga_banco.append(dict(chave=f['chave'], sha1=f['sha1']))
    page.locator('#cargaPasta').set_input_files(str(pasta_mc))      # sem alteração: nada é reenviado
    expect(page.locator('#cargaPasso')).to_contain_text('0 planilha(s) gravada(s) (0 itens), 1 sem alteração', timeout=30000)
    assert len(carga_enviada) == 1, 'reenviou planilha sem alteração'
    expect(page.locator('#btnCarga')).to_be_enabled()
    expect(page.locator('#sec-analise>summary')).to_contain_text('Gráficos')
    expect(page.locator('.disc-card')).to_have_count(3)
    expect(page.locator('.disc-card').nth(0)).to_contain_text('HH trabalhado e impacto ao longo do tempo')
    expect(page.locator('.disc-card').nth(1)).to_contain_text('HH trabalhado por disciplina')
    expect(page.locator('.disc-card').nth(2)).to_contain_text('HH de impacto por disciplina')
    expect(page.locator('.disc-echart-donut svg')).to_have_count(2)
    expect(page.locator('.disc-echart-line svg')).to_have_count(1)
    assert page.locator('.disc-echart-donut svg').first.locator('text').filter(has_text='%').count() >= 2
    assert page.evaluate("""() => {
      const linha = document.querySelector('.disc-wide').getBoundingClientRect();
      const rosca = document.querySelector('.graficos-disciplina>.disc-card:nth-child(2)').getBoundingClientRect();
      return linha.right < rosca.left && Math.abs(linha.top - rosca.top) < 2;
    }""")
    assert page.evaluate("""() => {
      const cols = getComputedStyle(document.querySelector('.graficos-disciplina')).gridTemplateColumns.split(' ');
      const rosca = document.querySelector('.graficos-disciplina>.disc-card:nth-child(2)');
      const plot = rosca.querySelector('.disc-donut-wrap').getBoundingClientRect();
      const legend = rosca.querySelector('.disc-legend').getBoundingClientRect();
      return cols.length === 12 && plot.width >= 200 && legend.left >= plot.right && legend.top < plot.bottom
        && getComputedStyle(rosca.querySelector('.disc-legend')).rowGap === '3px';
    }""")
    expect(page.locator('.disc-cycle-range')).to_have_text('Ciclo de medição · 21/09/2026 a 20/10/2026')
    assert any(c['p_ini'] == '2026-09-21' and c['p_fim'] == '2026-10-20' for c in export_calls)
    assert page.evaluate("""() => {
      const eixo = echarts.getInstanceByDom(document.querySelector('.disc-echart-line')).getOption().xAxis[0].data;
      return eixo.length === 30 && eixo[0] === '21/09' && eixo.at(-1) === '20/10'
        && cicloMedicao('2026-10-20').ini === '2026-09-21'
        && cicloMedicao('2026-10-21').fim === '2026-11-20'
        && cicloMedicao('2026-01-01').ini === '2025-12-21';
    }""")
    assert page.evaluate("""() => {
      const nodes = [...document.querySelectorAll('.disc-echart-donut, .disc-echart-line')];
      return nodes.length === 3 && nodes.every(node => {
        const chart = echarts.getInstanceByDom(node);
        return chart && chart.getOption().series.length > 0;
      });
    }""")
    assert page.evaluate("""() => {
      const plot = echarts.getInstanceByDom(document.querySelector('.disc-echart-line')).getOption();
      const css = getComputedStyle(document.documentElement);
      const tooltip = plot.tooltip[0].formatter([
        {dataIndex:10, value:2, seriesName:'Trabalhado', marker:'<span>●</span>'},
        {dataIndex:10, value:1, seriesName:'Impacto estimado', marker:'<span>◆</span>'}
      ]);
      return plot.yAxis[0].min === 0
        && plot.series[0].itemStyle.color === css.getPropertyValue('--chart-work').trim()
        && plot.series[1].itemStyle.color === css.getPropertyValue('--chart-impact').trim()
        && tooltip.includes('01/10/2026') && tooltip.includes('2 HH') && tooltip.includes('1 HH');
    }""")
    page.evaluate("""() => {
      const linhas = Array.from({length:9}, (_, i) => ({nome:`Disciplina ${i+1}`, hhMin:60, impactoMin:30}));
      const cores = PainelCharts.coresDisciplina(linhas.map(d => d.nome));
      PainelCharts.render(linhas, cores, [{data:'2026-10-01',nome:'Disciplina 1',hhMin:60,impactoMin:30}],
        '2026-10-01','2026-10-07',[], '2026-09-21','2026-10-20','pronto');
    }""")
    expect(page.locator('.disc-echart-bars svg')).to_have_count(2)
    expect(page.locator('.disc-card').nth(1).locator('.sr-only li')).to_have_count(9)
    assert page.evaluate("""() => {
      const card = document.querySelectorAll('.disc-card')[1];
      const chart = echarts.getInstanceByDom(card.querySelector('.disc-echart-bars'));
      const option = chart.getOption();
      const tip = option.tooltip[0].formatter({dataIndex:0});
      return option.series[0].type === 'bar' && option.xAxis[0].min === 0
        && option.yAxis[0].data.length === 9 && card.scrollWidth <= card.clientWidth + 1
        && tip.includes('Disciplina 9') && tip.includes('1 HH');
    }""")
    page.locator('#sec-analise').screenshot(path=str(OUT / 'graficos-9-disciplinas-1440.png'))
    page.set_viewport_size({'width': 390, 'height': 844})
    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    page.locator('#sec-analise').screenshot(path=str(OUT / 'graficos-9-disciplinas-390.png'))
    page.set_viewport_size({'width': 1440, 'height': 1000})
    page.evaluate('render()')
    expect(page.locator('.disc-card').nth(1).locator('.disc-legend li')).to_have_count(2)
    page.set_viewport_size({'width': 1920, 'height': 1000})
    page.evaluate("""() => {
      const linhas = ['PAC','REC','TEL/COB','ANDAIME'].map((nome, i) =>
        ({nome, hhMin:(i+1)*1800, impactoMin:(i+1)*240}));
      const eventos = Array.from({length:8}, (_, i) => ({data:`2026-10-${String(i+2).padStart(2,'0')}`,
        nome:'PAC', responsavel:'Teste', contrato:'A', contratoNome:'A', hhMin:(i+1)*600, impactoMin:i*60}));
      PainelCharts.render(linhas, PainelCharts.coresDisciplina(linhas.map(d => d.nome)), eventos,
        '2026-10-01','2026-10-10', eventos, '2026-09-21','2026-10-20','pronto');
    }""")
    expect(page.locator('.disc-data tbody tr')).to_have_count(8)
    assert page.evaluate("""() => [...document.querySelectorAll('.disc-card:not(.disc-wide)')].every(card => {
      const body = card.querySelector('.disc-donut-layout').getBoundingClientRect();
      const chart = card.querySelector('.disc-donut-wrap').getBoundingClientRect();
      return chart.height >= 240 && chart.height >= body.height * .75 && chart.width >= 240;
    })""")
    page.locator('#sec-analise').screenshot(path=str(OUT / 'graficos-8-dias-1920.png'))
    page.set_viewport_size({'width': 1440, 'height': 1000})
    page.evaluate('render()')
    expect(page.locator('.disc-data tbody tr')).to_have_count(3)
    assert page.evaluate("""() => {
      const rows = horasPorDisciplina(F.rdcs, dados.atividades.filter(a => F.rdcs.some(r => r.codigo === a.rdc)));
      return rows.length === 2 && rows.reduce((s, d) => s + d.hhMin, 0) === 360
        && rows.reduce((s, d) => s + d.impactoMin, 0) === 180;
    }""")
    assert page.evaluate("""() => {
      const eventos = eventosPorDisciplina(F.rdcs, dados.atividades);
      const dias = PainelCharts.periodizar(eventos, 'diaria', '2026-10-01', '2026-10-07');
      const semanas = PainelCharts.periodizar(eventos, 'semanal', '2026-10-01', '2026-10-07');
      const meses = PainelCharts.periodizar(eventos, 'mensal', '2026-10-01', '2026-10-07');
      return dias.length === 7 && dias.filter(d => d.registros).length === 3
        && dias[3].registros === 0 && dias[0].hhMin === 120
        && semanas.length === 2 && semanas[0].hhMin === 360
        && meses.length === 1 && meses[0].impactoMin === 180
        && PainelCharts.periodizar([{data:'2026-10-31',hhMin:60,impactoMin:30},
          {data:'2026-11-01',hhMin:120,impactoMin:60}], 'mensal', '2026-10-31', '2026-11-01')
          .map(d => d.hhMin).join(',') === '60,120';
    }""")
    assert page.evaluate("""() => {
      const chart = echarts.getInstanceByDom(document.querySelector('.disc-echart-line'));
      const series = chart.getOption().series;
      return series[0].data[3] == null && series[1].data[3] == null
        && series[0].connectNulls === false;
    }""")
    page.locator('.disc-period button').nth(1).click()
    expect(page.locator('.disc-period button').nth(1)).to_have_attribute('aria-pressed', 'true')
    expect(page.locator('.disc-data tbody tr')).to_have_count(1)
    page.locator('.disc-period button').nth(2).click()
    expect(page.locator('.disc-data tbody tr')).to_have_count(1)
    page.locator('.disc-chart-controls select').nth(0).select_option('PAC')
    expect(page.locator('.disc-data tbody tr td').first).to_have_text('2')
    expect(page.locator('.production-hero .metric-value')).to_have_text('6HH')
    expect(page.locator('.disc-donut-center strong').first).to_have_text('6')
    page.locator('.disc-chart-controls select').nth(0).select_option('')
    page.locator('.disc-chart-controls select').nth(2).select_option('__sem_contrato__')
    expect(page.locator('.disc-data tbody tr td').first).to_have_text('2')
    page.locator('.disc-chart-controls select').nth(2).select_option('')
    page.locator('.disc-chart-controls select').nth(1).select_option('Responsável de teste')
    expect(page.locator('.disc-data tbody tr td').first).to_have_text('6')
    page.locator('.disc-chart-controls select').nth(1).select_option('')
    page.evaluate("dados.rdcs[2].responsavel='Outro responsável'; render()")
    page.locator('.disc-chart-controls select').nth(1).select_option('Outro responsável')
    expect(page.locator('.disc-data tbody tr td').first).to_have_text('2')
    page.locator('.disc-chart-controls select').nth(1).select_option('Responsável de teste')
    expect(page.locator('.disc-data tbody tr td').first).to_have_text('4')
    page.locator('.disc-chart-controls select').nth(1).select_option('')
    page.locator('#btnAtualizar').click()
    page.locator('.disc-period button').nth(0).click()
    state['mode'] = 'cycle_error'
    page.locator('#btnAtualizar').click()
    expect(page.locator('#conteudo')).to_have_attribute('aria-busy', 'false')
    expect(page.locator('.disc-card').nth(0)).to_contain_text('Não foi possível carregar o ciclo de medição.')
    expect(page.locator('.disc-donut-center strong').first).to_have_text('6')
    page.locator('.disc-period button').nth(1).click()
    expect(page.locator('.disc-echart-line svg')).to_have_count(1)
    state['mode'] = 'ok'
    page.locator('#btnAtualizar').click()
    page.locator('.disc-period button').nth(0).click()
    expect(page.locator('.disc-echart-line svg')).to_have_count(1)
    assert page.evaluate("""() => {
      const color = selector => getComputedStyle(document.querySelector(selector)).color;
      return [color('#sec-analise>summary h2'), color('#tblResp thead th'), color('#tblResp tbody td.num')].join('|') ===
        ['rgb(40, 74, 104)', 'rgb(61, 92, 119)', 'rgb(82, 107, 130)'].join('|');
    }""")
    assert page.evaluate("""() => {
      const fill = selector => getComputedStyle(document.querySelector(selector)).backgroundColor;
      return [fill('#sec-analise>summary'), fill('#tblResp thead th'), fill('#tblResp tbody tr:nth-child(2) td')].join('|') ===
        ['rgb(229, 240, 248)', 'rgb(250, 251, 253)', 'rgb(250, 251, 253)'].join('|');
    }""")
    assert page.evaluate('planilhas().RDC.reduce((s,r)=>s+r.HH,0)') == 6
    page.evaluate("""() => {
      const hoje = new Date(), pad = n => String(n).padStart(2, '0');
      $('fIni').value = `${hoje.getFullYear()}-${pad(hoje.getMonth()+1)}-01`;
      $('fFim').value = `${hoje.getFullYear()}-${pad(hoje.getMonth()+1)}-${pad(hoje.getDate())}`;
    }""")
    page.locator('#btnAtualizar').click()
    expect(page.locator('#matriz .mat-stage.com-marcos')).to_have_count(1)
    remaining = page.evaluate('new Date(new Date().getFullYear(), new Date().getMonth()+1, 0).getDate() - new Date().getDate()')
    expect(page.locator('#matriz thead tr.dia th.futuro').first).to_be_visible()
    expect(page.locator('#btnFuturo')).to_be_visible()
    assert page.evaluate("""() => {
      const intro = document.querySelector('.emissao-intro').getBoundingClientRect();
      const button = document.querySelector('#btnFuturo').getBoundingClientRect();
      const table = document.querySelector('#matriz table.mat').getBoundingClientRect();
      return button.top < intro.bottom && table.top - intro.bottom <= 65;
    }""")
    assert page.locator('#matriz thead tr.dia th.futuro').count() == remaining
    expect(page.locator('#matriz .mat-line.hoje')).to_have_count(1)
    expect(page.locator('#matriz .mat-line.inicio-ciclo')).to_have_count(1)
    expect(page.locator('#matriz .mat-line.termino-ciclo')).to_have_count(1)
    expect(page.locator('#matriz .mat-line-label.hoje')).to_have_text('Hoje')
    expect(page.locator('#matriz .mat-line-label.inicio-ciclo')).to_have_text('Início')
    expect(page.locator('#matriz .mat-line-label.termino-ciclo')).to_have_text('Término')
    assert page.evaluate("""() => [...document.querySelectorAll('#matriz .mat-line-label')].every(e => getComputedStyle(e).writingMode === 'vertical-rl')""")
    assert page.evaluate("""() => [...document.querySelectorAll('#matriz .mat-line-label')].every(e => getComputedStyle(e).backgroundColor === 'rgba(0, 0, 0, 0)' && getComputedStyle(e).backgroundImage === 'none')""")
    today_line = page.locator('#matriz .mat-line.hoje')
    assert today_line.evaluate("e => getComputedStyle(e).backgroundImage.includes('repeating-linear-gradient')")
    assert today_line.evaluate("e => e.getBoundingClientRect().width == 2")
    assert page.evaluate("""() => [...document.querySelectorAll('#matriz .mat-line, #matriz .mat-key-line')].every(e => e.getBoundingClientRect().width === 2)""")
    assert today_line.evaluate("e => e.getBoundingClientRect().height > 300")
    assert page.evaluate("""() => {
      const th = document.querySelector('#matriz thead tr.dia th.hoje').getBoundingClientRect();
      const line = document.querySelector('#matriz .mat-line.hoje').getBoundingClientRect();
      return line.left >= th.left && line.right <= th.right;
    }""")
    assert page.evaluate("""() => [...document.querySelectorAll('#matriz .mat-line-label')].every(e => {
      const stage = e.parentElement.getBoundingClientRect(), r = e.getBoundingClientRect();
      return r.left >= stage.left && r.right <= stage.right && r.bottom < e.parentElement.querySelector('table').getBoundingClientRect().top;
    })""")
    assert page.locator('#matriz tbody tr').first.locator('td.d.futuro').count() == remaining
    assert page.evaluate('FERIADOS.size') == 28
    assert page.locator('#matriz tbody tr').first.locator('td.d.futuro .c.futuro').first.inner_text() == '–'
    feriado = page.locator('#matriz thead tr.dia th[title*="12/10/2026"]')
    assert 'feriado' in feriado.get_attribute('class').split()
    expect(feriado).to_have_attribute('title', '12/10/2026 · Nossa Senhora Aparecida · Nacional · data futura')
    assert page.locator('#matriz tbody tr').first.locator('td.d').nth(11).locator('.c').inner_text() == ''
    domingo = page.locator('#matriz thead tr.dia th[title*="11/10/2026"]')
    assert 'fds' in domingo.get_attribute('class').split()
    assert 'feriado' not in domingo.get_attribute('class').split()
    assert page.locator('#matriz tbody tr').first.locator('td.d').nth(10).locator('.c').inner_text() == ''
    expect(page.locator('.production-hero .metric-value')).to_have_text('6HH')
    page.locator('#sec-emissao').evaluate('(e)=>e.open=true')
    page.locator('#btnFuturo').click()
    expect(page.locator('#matriz thead tr.dia th.futuro').first).to_be_focused()
    page.locator('#sec-emissao').screenshot(path=str(OUT / 'emissao-dias-futuros.png'))
    page.set_viewport_size({'width': 390, 'height': 844})
    assert page.evaluate('document.documentElement.scrollWidth <= innerWidth')
    assert page.evaluate("document.querySelector('#matriz .mat-line.hoje').getBoundingClientRect().height > 300")
    assert page.evaluate("""() => {
      const th = document.querySelector('#matriz thead tr.dia th.hoje').getBoundingClientRect();
      const line = document.querySelector('#matriz .mat-line.hoje').getBoundingClientRect();
      return line.left >= th.left && line.right <= th.right;
    }""")
    assert page.evaluate("document.querySelector('#matriz tbody tr').getBoundingClientRect().height == 49")
    assert page.evaluate("document.querySelector('#matriz tbody td.nome').getBoundingClientRect().width <= 210")
    page.locator('#btnFuturo').click()
    page.wait_for_timeout(300)
    assert page.evaluate("document.querySelector('#matriz').scrollLeft > 0")
    page.locator('#sec-emissao').screenshot(path=str(OUT / 'emissao-mobile-marcos.png'))
    page.set_viewport_size({'width': 1440, 'height': 1000})
    page.evaluate("$('fIni').value='2026-10-20'; $('fFim').value='2026-11-21'")
    page.locator('#btnAtualizar').click()
    expect(page.locator('#matriz .mat-line.inicio-ciclo')).to_have_count(2)
    expect(page.locator('#matriz .mat-line.termino-ciclo')).to_have_count(2)
    expect(page.locator('#matriz thead tr.dia th.inicio-ciclo').first).to_have_attribute('title', '21/10/2026 · Início do ciclo de medição · 21/10/2026 a 20/11/2026 · data futura')
    expect(page.locator('#matriz thead tr.dia th.termino-ciclo').last).to_have_attribute('title', '20/11/2026 · Término da medição · 21/10/2026 a 20/11/2026 · Dia Nacional da Consciência Negra · Nacional · data futura')
    page.evaluate("$('fIni').value='2026-12-20'; $('fFim').value='2027-01-21'")
    page.locator('#btnAtualizar').click()
    expect(page.locator('#matriz thead tr.dia th.termino-ciclo').last).to_have_attribute('title', '20/01/2027 · Término da medição · 21/12/2026 a 20/01/2027 · data futura')
    page.evaluate("$('fIni').value='2026-10-01'; $('fFim').value='2026-10-07'")
    page.locator('#btnAtualizar').click()
    expect(page.locator('#btnFuturo')).to_be_hidden()
    page.evaluate("$('fIni').value='2026-09-07'; $('fFim').value='2026-09-09'")
    page.locator('#btnAtualizar').click()
    expect(page.locator('#matriz thead tr.dia th.feriado')).to_have_count(2)
    expect(page.locator('#matriz thead tr.dia th.feriado').nth(1)).to_have_attribute('title', '08/09/2026 · Aniversário de São Luís · Municipal (São Luís)')
    assert page.locator('#matriz tbody tr').first.locator('.resp-texto small').count() == 0
    page.evaluate("$('fIni').value='2027-05-01'; $('fFim').value='2027-05-01'")
    page.locator('#btnAtualizar').click()
    expect(page.locator('#matriz thead tr.dia th.feriado')).to_have_count(1)
    assert {'feriado', 'fds'} <= set(page.locator('#matriz thead tr.dia th.feriado').get_attribute('class').split())
    assert page.locator('#matriz tbody tr').first.locator('td.d .c').inner_text() == ''
    page.evaluate("dados.rdcs[0].data_rdc='2027-05-01'; render()")
    assert page.locator('#matriz tbody tr').first.locator('td.d .c').inner_text() == ''
    assert page.locator('#matriz tbody tr').last.locator('td.d .c').inner_text() == '✓'
    page.evaluate("$('fIni').value='2026-10-01'; $('fFim').value='2026-10-07'")
    page.locator('#btnAtualizar').click()
    expect(page.locator('#matriz thead tr.dia th')).to_have_count(8)
    for width in [1440, 1024, 768, 390, 320]:
        page.set_viewport_size({'width': width, 'height': 1000})
        page.evaluate('document.fonts.ready')
        page.wait_for_timeout(150)
        page.evaluate('document.activeElement.blur(); scrollTo({top:0,behavior:"instant"})')
        assert page.evaluate('document.documentElement.scrollWidth <= innerWidth'), width
        assert page.evaluate("""() => [...document.querySelectorAll('.disc-card:not(.disc-wide)')].every(card => {
          const legend = card.querySelector('.disc-legend');
          if (!legend) return true;
          const outer = card.getBoundingClientRect(), inner = legend.getBoundingClientRect();
          return inner.left >= outer.left && inner.right <= outer.right + 1 && legend.scrollWidth <= legend.clientWidth + 1;
        })"""), width
        assert page.evaluate("getComputedStyle(document.querySelector('#kpis')).gridTemplateColumns.split(' ').length === 12"), width
        assert page.evaluate("[...document.querySelectorAll('#kpis .metric')].every(e=>e.scrollWidth<=e.clientWidth)"), (
            width, page.evaluate("[...document.querySelectorAll('#kpis .metric')].map(e=>({label:e.innerText.split('\\n')[0],scroll:e.scrollWidth,client:e.clientWidth}))")
        )
        if width <= 390:
            assert page.evaluate("""() => {
              const labels = [...document.querySelectorAll('.disc-chart-controls label')];
              return labels.every((label, i) => !i || label.getBoundingClientRect().top > labels[i - 1].getBoundingClientRect().bottom);
            }"""), width
        page.screenshot(path=str(OUT / f'painel-{width}.png'), full_page=True)

    page.set_viewport_size({'width': 1440, 'height': 1000})
    page.locator('#btnPendencias').click()
    expect(page.locator('#fStatus')).to_have_value('Enviado')
    expect(page.locator('#tblRdc tbody tr[data-r]')).to_have_count(1)
    header = page.locator('#tblRdc th[data-c="1"]')
    header.focus()
    page.keyboard.press('Enter')
    expect(header).to_have_attribute('aria-sort', 'ascending')
    page.locator('#tblRdc tr[data-r]').focus()
    page.keyboard.press('Enter')
    expect(page.locator('#modal')).to_be_visible()
    page.keyboard.press('Escape')

    page.locator('#btnLimpar').click()
    page.locator('#fDetalhes').evaluate('(e)=>e.open=true')
    page.locator('#fDisc').select_option('PAC')
    expect(page.locator('.production-hero .metric-value')).to_have_text('2HH')
    expect(page.locator('#dataContext')).to_be_hidden()
    expect(page.locator('#kpis .metric.impact')).to_have_attribute('title', 'Impacto rateado proporcionalmente ao HH trabalhado da disciplina.')
    expect(page.locator('#kpis .metric.impact .metric-value')).to_have_text('1HH')
    assert page.evaluate("horasPorDisciplina(F.rdcs, dados.atividades).length === 1")
    page.locator('#fDisc').select_option('REC')
    expect(page.locator('#kpis .metric').nth(4)).to_have_attribute('title', 'Horários não informados')
    page.locator('#btnLimpar').click()
    assert page.evaluate("""() => {
      const novo = {...dados.atividades[0], codigo:'A-extra', disciplina:'REC', duracao_min:60};
      dados.atividades.push(novo);
      render();
      const rows = horasPorDisciplina(F.rdcs, dados.atividades);
      const pac = rows.find(d => d.nome === 'PAC'), rec = rows.find(d => d.nome === 'REC');
      return rows.reduce((s,d)=>s+d.impactoMin,0) === 180
        && pac.impactoMin === 30 && rec.impactoMin === 150
        && rows.reduce((s,d)=>s+d.hhMin,0) === 480;
    }""")
    page.locator('#fDisc').select_option('PAC')
    assert page.evaluate("""() => {
      const chart = echarts.getInstanceByDom(document.querySelector('.resp-echart')).getOption();
      const impactoTabela = document.querySelector('#tblResp tbody tr:last-child td:nth-child(7)').textContent.trim();
      return chart.series[0].data[0] === 2
        && chart.series[1].data[0] === .5
        && impactoTabela === '0,5'
        && document.querySelector('#kpis .metric.impact .metric-value').textContent.trim() === '0,5HH';
    }""")
    page.locator('#btnLimpar').click()
    page.locator('#btnAtualizar').click()
    expect(page.locator('#conteudo')).to_have_attribute('aria-busy', 'false')

    page.evaluate("dados.rdcs.forEach(r => r.pts=false); render()")
    expect(page.locator('#kpis .metric').nth(4)).to_have_attribute('title', 'Nenhum RDC com PTS')
    page.locator('#btnAtualizar').click()
    expect(page.locator('#kpis .metric').nth(4)).to_have_attribute('title', 'Média · 1 de 2 RDC com PTS')

    for mode in ['error', 'network']:
        state['mode'] = mode
        page.locator('#btnAtualizar').click()
        expect(page.locator('#dataContext')).to_contain_text('última consulta carregada')
        expect(page.locator('#conteudo')).to_have_attribute('aria-busy', 'false')
    state['mode'] = 'empty'
    page.locator('#btnAtualizar').click()
    expect(page.locator('#vazio')).to_be_visible()
    expect(page.locator('#sec-analise')).to_be_hidden()
    expect(page.locator('#resumoExecutivo')).to_be_hidden()
    expect(page.locator('#dataContext')).not_to_contain_text('última consulta carregada')
    state['mode'] = 'ok'
    page.locator('#btnAtualizar').click()
    expect(page.locator('#vazio')).to_be_hidden()
    page.evaluate("$('fIni').value='2026-01-01'; $('fFim').value='2026-12-31'")
    page.locator('#btnAtualizar').click()
    expect(page.locator('.disc-card')).to_have_count(3)
    assert not errors, errors
    print('PASS: login, gráficos por disciplina, rateio de impacto, KPIs, recortes, exportação, teclado, erros, vazio, recuperação e 5 larguras.')
    print(f'Capturas: {OUT}')
    browser.close()
